#!/usr/bin/env python3
"""Check CSMAR zip files for a minimally complete Excel payload.

Usage: python verify_integrity.py <data-directory>

The heuristic is intentionally conservative and CSMAR-oriented: more than
three columns is reported as OK, while three or fewer columns are flagged for
manual review. It is not a substitute for checking the requested schema.
"""

from __future__ import annotations

import json
import os
import shutil
import sys
import tempfile
import warnings
import zipfile

from openpyxl import load_workbook

warnings.filterwarnings("ignore")


def verify_zip(zip_path: str) -> dict:
    size_mb = round(os.path.getsize(zip_path) / 1024 / 1024, 1)
    fields = 0
    rows = 0
    headers: list[str] = []

    try:
        with zipfile.ZipFile(zip_path, "r") as archive:
            xlsx_files = [name for name in archive.namelist() if name.lower().endswith(".xlsx")]
            if not xlsx_files:
                return {"size_mb": size_mb, "fields": 0, "rows": 0, "status": "NONE", "headers": []}

            with tempfile.TemporaryDirectory() as tmpdir:
                archive.extract(xlsx_files[0], tmpdir)
                workbook_path = os.path.join(tmpdir, xlsx_files[0])
                workbook = load_workbook(workbook_path, read_only=False, data_only=False)
                try:
                    sheet = workbook.active
                    fields = sheet.max_column or 0
                    rows = max((sheet.max_row or 1) - 1, 0)
                    for column in range(1, min(6, fields + 1)):
                        value = sheet.cell(row=1, column=column).value
                        if value is not None and str(value).strip():
                            headers.append(str(value))
                finally:
                    workbook.close()

        status = "OK" if fields > 3 else "BAD"
    except Exception as exc:  # Report the file and continue scanning the batch.
        return {
            "size_mb": size_mb,
            "fields": -1,
            "rows": 0,
            "status": "ERR",
            "headers": [],
            "error": str(exc)[:200],
        }

    return {"size_mb": size_mb, "fields": fields, "rows": rows, "status": status, "headers": headers}


def main() -> int:
    if len(sys.argv) < 2:
        print("Usage: python verify_integrity.py <data-directory>")
        return 1

    base_dir = sys.argv[1]
    if not os.path.isdir(base_dir):
        print(f"Directory not found: {base_dir}")
        return 1

    results: list[dict] = []
    for database_name in sorted(os.listdir(base_dir)):
        database_path = os.path.join(base_dir, database_name)
        if not os.path.isdir(database_path):
            continue
        for filename in sorted(os.listdir(database_path)):
            if not filename.lower().endswith(".zip"):
                continue
            record = verify_zip(os.path.join(database_path, filename))
            record.update({"db": database_name, "file": filename})
            results.append(record)

    report_path = os.path.join(base_dir, "integrity_report.json")
    with open(report_path, "w", encoding="utf-8") as report:
        json.dump(results, report, ensure_ascii=False, indent=2)

    ok_files = [item for item in results if item["status"] == "OK"]
    review_files = [item for item in results if item["status"] != "OK"]
    print("CSMAR integrity report")
    print(f"Total: {len(results)} | OK: {len(ok_files)} | Review: {len(review_files)}")
    for item in review_files:
        print(f"[{item['status']}] {item['db']}/{item['file']} fields={item['fields']}")
    print(f"Detailed report: {report_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
