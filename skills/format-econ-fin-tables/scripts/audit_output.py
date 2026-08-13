#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path


def audit_tex(path: Path):
    text = path.read_text(encoding="utf-8")
    errors = []
    warnings = []
    for token in (r"\toprule", r"\midrule", r"\bottomrule"):
        if token not in text:
            errors.append(f"missing {token}")
    if r"\hline" in text:
        errors.append("uses \\hline instead of booktabs rules")
    if r"\resizebox" in text or r"\scalebox" in text:
        errors.append("whole-table scaling is not allowed by default")
    for match in re.finditer(r"\\begin\{(?:tabular|longtable)\}\{([^}]*)\}", text):
        if "|" in match.group(1):
            errors.append("vertical rule found in a column specification")
    if r"\begin{longtable}" in text:
        for token in (r"\endfirsthead", r"\endhead", r"\endfoot", r"\endlastfoot"):
            if token not in text:
                errors.append(f"longtable is missing {token}")
        if "continued" not in text.lower() and "（续）" not in text:
            errors.append("longtable is missing a continuation label")
    return errors, warnings


def audit_docx(path: Path):
    try:
        from docx import Document
        from docx.oxml.ns import qn
    except ImportError as exc:
        raise RuntimeError("python-docx is required to audit Word output") from exc
    doc = Document(path)
    errors = []
    warnings = []
    if not doc.tables:
        errors.append("document contains no tables")
        return errors, warnings
    continuation_titles = 0
    for paragraph in doc.paragraphs:
        text = paragraph.text.lower()
        if "(continued)" in text or "（续）" in paragraph.text:
            continuation_titles += 1
    for table_index, table in enumerate(doc.tables, 1):
        layout = table._tbl.tblPr.find(qn("w:tblLayout"))
        if layout is None or layout.get(qn("w:type")) != "fixed":
            errors.append(f"table {table_index} does not use fixed layout")
        header_rows = 0
        for row in table.rows:
            if row._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None:
                header_rows += 1
            for cell in row.cells:
                borders = cell._tc.get_or_add_tcPr().find(qn("w:tcBorders"))
                if borders is None:
                    continue
                for edge in ("left", "right", "insideV"):
                    element = borders.find(qn(f"w:{edge}"))
                    if element is not None and element.get(qn("w:val")) not in {None, "nil", "none"}:
                        errors.append(f"table {table_index} contains a vertical cell border")
        if header_rows == 0:
            errors.append(f"table {table_index} has no repeating header rows")
    if len(doc.tables) > 1 and continuation_titles < len(doc.tables) - 1:
        errors.append("segmented Word table is missing continuation titles")
    note_markers = ("Notes:", "Note:", "注：")
    if not any(any(marker in paragraph.text for marker in note_markers) for paragraph in doc.paragraphs):
        warnings.append("no table-note paragraph detected")
    return errors, warnings


def main():
    parser = argparse.ArgumentParser(description="Audit generated LaTeX or Word table structure.")
    parser.add_argument("output")
    args = parser.parse_args()
    path = Path(args.output)
    if not path.exists():
        print(f"ERROR: file not found: {path}")
        return 2
    try:
        if path.suffix.lower() == ".tex":
            errors, warnings = audit_tex(path)
        elif path.suffix.lower() == ".docx":
            errors, warnings = audit_docx(path)
        else:
            print("ERROR: audit_output.py supports .tex and .docx")
            return 2
    except Exception as exc:
        print(f"ERROR: {exc}")
        return 2
    for warning in warnings:
        print(f"WARNING: {warning}")
    for error in errors:
        print(f"ERROR: {error}")
    if not errors:
        print(f"OK: audited {path}")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
