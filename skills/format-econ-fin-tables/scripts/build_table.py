#!/usr/bin/env python3
from __future__ import annotations

import argparse
import sys
from pathlib import Path

from spec_utils import (
    contiguous_groups,
    display_cell,
    latex_escape,
    load_specs,
    safe_tex_label,
    split_logical_rows,
    validate_spec,
)


def header_rows_latex(spec):
    columns = spec["columns"]
    lines = []
    groups = contiguous_groups(columns)
    if any(name for name, _count in groups):
        cells = [""] + [rf"\multicolumn{{{count}}}{{c}}{{{latex_escape(name)}}}" for name, count in groups]
        lines.append(" & ".join(cells) + r" \\")
        cursor = 2
        rules = []
        for _name, count in groups:
            if count:
                rules.append(rf"\cmidrule(lr){{{cursor}-{cursor + count - 1}}}")
            cursor += count
        lines.append(" ".join(rules))
    if any(col.get("model") for col in columns):
        lines.append(" & ".join([""] + [latex_escape(col.get("model", "")) for col in columns]) + r" \\")
    variable_label = "变量" if spec.get("language") == "zh" else "Variables"
    lines.append(" & ".join([variable_label] + [latex_escape(col.get("label", "")) for col in columns]) + r" \\")
    lines.append(r"\midrule")
    return lines


def data_rows_latex(spec):
    column_keys = [col["key"] for col in spec["columns"]]
    lines = []
    for row in spec["rows"]:
        kind = row.get("kind")
        if kind == "spacer":
            lines.append(r"\addlinespace")
            continue
        if kind == "panel":
            lines.append(rf"\addlinespace\multicolumn{{{len(column_keys) + 1}}}{{l}}{{\textbf{{{latex_escape(row.get('label', ''))}}}}} \\")
            continue
        values = row.get("values") or {}
        cells = [latex_escape(row.get("label", ""))]
        cells.extend(latex_escape(display_cell(values.get(key))) for key in column_keys)
        suffix = r" \\*" if kind == "coefficient" else r" \\"
        lines.append(" & ".join(cells) + suffix)
    return lines


def build_latex(spec):
    columns = spec["columns"]
    ncols = len(columns) + 1
    colspec = "l" + "c" * len(columns)
    title = latex_escape(spec["title"])
    label = safe_tex_label(spec)
    notes = [latex_escape(note) for note in spec.get("notes", [])]
    long_table = bool(spec.get("multipage")) or len(spec["rows"]) > 25
    headers = header_rows_latex(spec)
    body = data_rows_latex(spec)
    if long_table:
        continuation_title = "（续）" if spec.get("language") == "zh" else " (continued)"
        continuation_footer = "续下页" if spec.get("language") == "zh" else "Continued on next page"
        result = [
            "% Requires booktabs, longtable, and threeparttablex.",
            r"\begin{ThreePartTable}",
            r"\begin{TableNotes}[flushleft]",
        ]
        result.extend(rf"\item {note}" for note in notes)
        result.extend(
            [
                r"\end{TableNotes}",
                rf"\begin{{longtable}}{{{colspec}}}",
                rf"\caption{{{title}}}\label{{{label}}} \\",
                r"\toprule",
                *headers,
                r"\endfirsthead",
                rf"\multicolumn{{{ncols}}}{{l}}{{\textbf{{{title}{continuation_title}}}}} \\",
                r"\toprule",
                *headers,
                r"\endhead",
                rf"\midrule\multicolumn{{{ncols}}}{{r}}{{{continuation_footer}}} \\",
                r"\endfoot",
                r"\bottomrule",
                r"\insertTableNotes",
                r"\endlastfoot",
                *body,
                r"\end{longtable}",
                r"\end{ThreePartTable}",
            ]
        )
        return "\n".join(result) + "\n"
    result = [
        "% Requires booktabs and threeparttable.",
        r"\begin{table}[!htbp]",
        r"\centering",
        rf"\caption{{{title}}}",
        rf"\label{{{label}}}",
        r"\begin{threeparttable}",
        rf"\begin{{tabular}}{{{colspec}}}",
        r"\toprule",
        *headers,
        *body,
        r"\bottomrule",
        r"\end{tabular}",
    ]
    if notes:
        result.append(r"\begin{tablenotes}[flushleft]\footnotesize")
        result.extend(rf"\item {note}" for note in notes)
        result.append(r"\end{tablenotes}")
    result.extend([r"\end{threeparttable}", r"\end{table}"])
    return "\n".join(result) + "\n"


def set_cell_border(cell, **edges):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, settings in edges.items():
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in settings.items():
            element.set(qn(f"w:{key}"), str(value))


def set_repeat_table_header(row):
    from docx.oxml import OxmlElement

    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "true")
    tr_pr.append(repeat)


def prevent_row_split(row):
    from docx.oxml import OxmlElement

    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_cell_text(cell, text, *, bold=False, italic=False, align="center", size=10, language="en"):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt

    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if align == "left" else WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.italic = italic
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体" if language == "zh" else "Times New Roman")
    cell.vertical_alignment = 1
    return paragraph


def remove_all_table_borders(table):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)


def set_fixed_table_layout(table, widths):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    indentation = tbl_pr.first_child_found_in("w:tblInd")
    if indentation is None:
        indentation = OxmlElement("w:tblInd")
        tbl_pr.append(indentation)
    indentation.set(qn("w:w"), "0")
    indentation.set(qn("w:type"), "dxa")
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = widths[index]
            tc_w = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_w.set(qn("w:w"), str(int(widths[index] / 635)))
            tc_w.set(qn("w:type"), "dxa")


def add_header_rows(table, spec, size):
    columns = spec["columns"]
    language = spec.get("language", "en")
    header_rows = []
    groups = contiguous_groups(columns)
    if any(name for name, _count in groups):
        row = table.add_row()
        set_cell_text(row.cells[0], "", size=size, language=language)
        cursor = 1
        for name, count in groups:
            merged = row.cells[cursor]
            if count > 1:
                merged = merged.merge(row.cells[cursor + count - 1])
            set_cell_text(merged, name, bold=True, size=size, language=language)
            cursor += count
        header_rows.append(row)
    if any(col.get("model") for col in columns):
        row = table.add_row()
        set_cell_text(row.cells[0], "", size=size, language=language)
        for index, col in enumerate(columns, 1):
            set_cell_text(row.cells[index], col.get("model", ""), size=size, language=language)
        header_rows.append(row)
    row = table.add_row()
    set_cell_text(row.cells[0], "变量" if language == "zh" else "Variables", bold=True, align="left", size=size, language=language)
    for index, col in enumerate(columns, 1):
        set_cell_text(row.cells[index], col.get("label", ""), bold=True, size=size, language=language)
    header_rows.append(row)
    for header in header_rows:
        set_repeat_table_header(header)
        prevent_row_split(header)
    for cell in header_rows[0].cells:
        set_cell_border(cell, top={"val": "single", "sz": "12", "color": "000000"})
    for cell in header_rows[-1].cells:
        set_cell_border(cell, bottom={"val": "single", "sz": "6", "color": "000000"})
    return header_rows


def add_data_rows(table, spec, rows, size):
    columns = spec["columns"]
    language = spec.get("language", "en")
    created = []
    for row_spec in rows:
        kind = row_spec.get("kind")
        row = table.add_row()
        prevent_row_split(row)
        if kind == "spacer":
            for cell in row.cells:
                set_cell_text(cell, "", size=max(size - 2, 7), language=language)
            row.height = 80000
            created.append(row)
            continue
        if kind == "panel":
            merged = row.cells[0].merge(row.cells[-1])
            set_cell_text(merged, row_spec.get("label", ""), bold=True, align="left", size=size, language=language)
            set_cell_border(merged, top={"val": "single", "sz": "4", "color": "000000"})
            created.append(row)
            continue
        label = row_spec.get("label", "")
        set_cell_text(row.cells[0], label, italic=kind in {"coefficient", "statistic"} and bool(label), align="left", size=size, language=language)
        values = row_spec.get("values") or {}
        for index, col in enumerate(columns, 1):
            set_cell_text(row.cells[index], display_cell(values.get(col["key"])), size=size, language=language)
        if kind == "coefficient":
            for cell in row.cells:
                cell.paragraphs[0].paragraph_format.keep_with_next = True
        created.append(row)
    return created


def build_word(spec, output):
    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt
    except ImportError as exc:
        raise RuntimeError("python-docx is required for Word output") from exc

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    if len(spec["columns"]) > 5:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10 if spec.get("language") == "en" else 9)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体" if spec.get("language") == "zh" else "Times New Roman")
    size = 10 if spec.get("language") == "en" else 9
    rows_per_page = int(spec.get("rows_per_page", 24))
    segments = split_logical_rows(spec["rows"], rows_per_page)
    usable = section.page_width - section.left_margin - section.right_margin
    label_width = int(usable * (0.28 if spec["table_type"] == "variable_definition" else 0.24))
    result_width = int((usable - label_width) / len(spec["columns"]))
    widths = [label_width] + [result_width] * len(spec["columns"])
    for segment_index, segment in enumerate(segments):
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_after = Pt(4)
        continued = segment_index > 0
        continuation = "（续）" if spec.get("language") == "zh" else " (continued)"
        label_text = f"{spec['label']}{continuation if continued else ''}"
        run = title.add_run(f"{label_text}  {spec['title']}")
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(size + 0.5)
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体" if spec.get("language") == "zh" else "Times New Roman")
        if segment_index == 0 and spec.get("caption_note"):
            paragraph = doc.add_paragraph(str(spec["caption_note"]))
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.keep_with_next = True
        table = doc.add_table(rows=0, cols=len(spec["columns"]) + 1)
        remove_all_table_borders(table)
        add_header_rows(table, spec, size)
        data_rows = add_data_rows(table, spec, segment, size)
        set_fixed_table_layout(table, widths)
        if data_rows:
            for cell in data_rows[-1].cells:
                set_cell_border(cell, bottom={"val": "single", "sz": "12", "color": "000000"})
        if segment_index < len(segments) - 1:
            doc.add_page_break()
    notes = spec.get("notes") or []
    if notes:
        note = doc.add_paragraph()
        note.paragraph_format.space_before = Pt(4)
        note.paragraph_format.space_after = Pt(0)
        prefix = "注：" if spec.get("language") == "zh" else "Notes: "
        run = note.add_run(prefix + " ".join(str(item) for item in notes))
        run.font.name = "Times New Roman"
        run.font.size = Pt(max(size - 1, 8))
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体" if spec.get("language") == "zh" else "Times New Roman")
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main():
    parser = argparse.ArgumentParser(description="Build LaTeX or editable Word economics/finance tables.")
    parser.add_argument("spec")
    parser.add_argument("--format", choices=("latex", "word"), required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--force", action="store_true", help="Allow replacement of an existing output")
    args = parser.parse_args()
    output = Path(args.output)
    if output.exists() and not args.force:
        print(f"ERROR: output exists: {output}; use --force to replace", file=sys.stderr)
        return 2
    specs = load_specs(args.spec)
    if len(specs) != 1:
        print("ERROR: build_table.py accepts exactly one table specification", file=sys.stderr)
        return 2
    spec = specs[0]
    errors, warnings = validate_spec(spec)
    for warning in warnings:
        print(f"WARNING: {warning}", file=sys.stderr)
    if errors:
        for error in errors:
            print(f"ERROR: {error}", file=sys.stderr)
        return 1
    output.parent.mkdir(parents=True, exist_ok=True)
    if args.format == "latex":
        output.write_text(build_latex(spec), encoding="utf-8")
    else:
        build_word(spec, output)
    print(f"OK: wrote {output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
